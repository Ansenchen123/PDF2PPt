from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Pt

from pdf2ppt.geometry import bbox_to_slide_units
from pdf2ppt.models import SlideLayout, TextBlock


ALIGNMENT_MAP = {
    "left": PP_ALIGN.LEFT,
    "center": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT,
    "justify": PP_ALIGN.JUSTIFY,
}


def create_pptx(layouts: list[SlideLayout], output_ppt_path: str | Path) -> None:
    if not layouts:
        raise ValueError("At least one slide layout is required")

    output_ppt_path = Path(output_ppt_path)
    output_ppt_path.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Pt(layouts[0].width_pt)
    prs.slide_height = Pt(layouts[0].height_pt)
    blank_layout = prs.slide_layouts[6]

    for layout in layouts:
        slide = prs.slides.add_slide(blank_layout)
        background = layout.clean_background_path or layout.background_image_path
        slide.shapes.add_picture(
            str(background),
            0,
            0,
            width=prs.slide_width,
            height=prs.slide_height,
        )
        for block in layout.text_blocks:
            _add_text_block(slide, block, prs.slide_width, prs.slide_height)

    prs.save(str(output_ppt_path))


def _add_text_block(slide, block: TextBlock, slide_width: int, slide_height: int) -> None:
    left, top, width, height = bbox_to_slide_units(block.bbox, slide_width, slide_height)
    box = slide.shapes.add_textbox(left, top, width, height)
    text_frame = box.text_frame
    text_frame.clear()
    text_frame.word_wrap = True
    text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    text_frame.margin_left = 0
    text_frame.margin_right = 0
    text_frame.margin_top = 0
    text_frame.margin_bottom = 0

    paragraph = text_frame.paragraphs[0]
    paragraph.alignment = ALIGNMENT_MAP.get(block.style.alignment, PP_ALIGN.LEFT)
    paragraph.line_spacing = block.style.line_spacing
    run = paragraph.add_run()
    run.text = block.text
    run.font.name = block.style.font_family or "Aptos"
    run.font.size = Pt(block.style.font_size_pt)
    run.font.bold = block.style.bold
    run.font.italic = block.style.italic
    run.font.color.rgb = _hex_to_rgb(block.style.color_hex)


def _hex_to_rgb(value: str) -> RGBColor:
    value = value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))
