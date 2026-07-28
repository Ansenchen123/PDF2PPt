from pathlib import Path

import fitz
from PIL import Image, ImageDraw
from pptx import Presentation

from pdf2ppt.background import clean_background
from pdf2ppt.models import BoundingBox, SlideLayout, TextBlock, TextStyle
from pdf2ppt.pdf import extract_native_text_blocks, render_pdf_pages
from pdf2ppt.pipeline import convert_pdf_to_ppt
from pdf2ppt.pptx_export import create_pptx
from pdf2ppt.reconstruction import build_slide_layout


def _make_text_pdf(path: Path) -> None:
    doc = fitz.open()
    page = doc.new_page(width=800, height=450)
    page.insert_text((80, 120), "Quarterly Revenue", fontsize=32, color=(0.1, 0.2, 0.3))
    doc.save(path)
    doc.close()


def _make_background(path: Path, size: tuple[int, int] = (800, 450)) -> None:
    img = Image.new("RGB", size, "#f8fafc")
    draw = ImageDraw.Draw(img)
    draw.rectangle((40, 40, size[0] - 40, size[1] - 40), outline="#cbd5e1", width=3)
    draw.text((80, 120), "Quarterly Revenue", fill="#111827")
    img.save(path)


def test_render_pdf_pages_preserves_page_dimensions(tmp_path: Path):
    pdf_path = tmp_path / "sample.pdf"
    _make_text_pdf(pdf_path)

    pages = render_pdf_pages(pdf_path, tmp_path / "pages", start_page=1, end_page=1, zoom=1)

    assert len(pages) == 1
    assert pages[0].page_width_pt == 800
    assert pages[0].page_height_pt == 450
    assert pages[0].image_path.exists()


def test_extract_native_text_blocks_keeps_style_and_bbox(tmp_path: Path):
    pdf_path = tmp_path / "sample.pdf"
    _make_text_pdf(pdf_path)

    blocks = extract_native_text_blocks(pdf_path, page_number=1)

    assert blocks
    assert blocks[0].text == "Quarterly Revenue"
    assert blocks[0].source == "native_pdf"
    assert blocks[0].style.font_size_pt == 32
    assert blocks[0].style.color_hex.startswith("#")


def test_clean_background_creates_mask_and_clean_image(tmp_path: Path):
    image_path = tmp_path / "page.png"
    clean_path = tmp_path / "clean.png"
    mask_path = tmp_path / "mask.png"
    overlay_path = tmp_path / "overlay.png"
    _make_background(image_path)
    block = TextBlock(
        id="title",
        text="Quarterly Revenue",
        bbox=BoundingBox(y_min=240, x_min=90, y_max=330, x_max=520),
        style=TextStyle(font_size_pt=32),
    )

    result = clean_background(image_path, [block], clean_path, mask_path, overlay_path)

    assert result.clean_image_path.exists()
    assert result.mask_path.exists()
    assert result.overlay_path.exists()
    assert result.mask_pixels > 0


def test_create_pptx_writes_slide_size_and_text_style(tmp_path: Path):
    background = tmp_path / "bg.png"
    output = tmp_path / "deck.pptx"
    _make_background(background)
    layout = SlideLayout(
        page_number=1,
        width_pt=800,
        height_pt=450,
        background_image_path=background,
        clean_background_path=background,
        text_blocks=[
            TextBlock(
                id="title",
                text="Quarterly Revenue",
                bbox=BoundingBox(y_min=240, x_min=90, y_max=330, x_max=520),
                style=TextStyle(font_size_pt=32, color_hex="#123456", bold=True),
            )
        ],
    )

    create_pptx([layout], output)

    prs = Presentation(output)
    assert round(prs.slide_width.pt) == 800
    assert round(prs.slide_height.pt) == 450
    text_shapes = [shape for shape in prs.slides[0].shapes if shape.has_text_frame and shape.text]
    assert len(text_shapes) == 1
    run = text_shapes[0].text_frame.paragraphs[0].runs[0]
    assert run.font.size.pt == 32
    assert run.font.bold is True


def test_pipeline_with_mock_provider_creates_pptx(tmp_path: Path):
    pdf_path = tmp_path / "sample.pdf"
    output = tmp_path / "sample.pptx"
    _make_text_pdf(pdf_path)

    result = convert_pdf_to_ppt(
        pdf_path=pdf_path,
        output_ppt=output,
        provider_name="mock",
        start_page=1,
        end_page=1,
        keep_temp=True,
    )

    assert output.exists()
    assert result.output_ppt == output
    assert len(result.layouts) == 1
    assert result.layouts[0].text_blocks


def test_build_slide_layout_uses_provider_result_for_image_pages(tmp_path: Path):
    image_path = tmp_path / "page.png"
    _make_background(image_path)
    block = TextBlock(id="a", text="A", bbox=BoundingBox(y_min=100, x_min=100, y_max=200, x_max=300))

    layout = build_slide_layout(
        page_number=1,
        width_pt=800,
        height_pt=450,
        image_path=image_path,
        provider_blocks=[block],
        native_blocks=[],
    )

    assert layout.text_blocks == [block]
    assert layout.background_image_path == image_path
